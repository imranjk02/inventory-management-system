from flask import Flask, render_template, request, redirect, url_for, session, flash
from database import get_db_connection
from flask import send_file
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from io import BytesIO
import os
import barcode
from barcode.writer import ImageWriter

app = Flask(__name__)
app.secret_key = "nettech_inventory_secret"

# LOGIN
@app.route("/", methods=["GET", "POST"])
def login():

    if request.method == "POST":
        username = request.form["username"]
        password = request.form["password"]
        connection = get_db_connection()

        if connection is None:
            flash("Database connection failed.", "error")
            return redirect(url_for("login"))

        cursor = connection.cursor(dictionary=True)
        query = """SELECT * FROM admins WHERE username = %s AND password = %s"""
        cursor.execute(query,(username, password))
        admin = cursor.fetchone()
        cursor.close()
        connection.close()

        if admin:
            session["admin_id"] = admin["id"]
            session["username"] = admin["username"]
            return redirect(url_for("dashboard"))
        else:
            flash("Invalid username or password.", "error")
    return render_template("login.html")


# DASHBOARD
@app.route("/dashboard")
def dashboard():

    if "admin_id" not in session:
        return redirect(url_for("login"))

    connection = get_db_connection()
    cursor = connection.cursor(dictionary=True)

    # Total courses
    cursor.execute("SELECT COUNT(*) AS total FROM courses")
    total_courses = cursor.fetchone()["total"]

    # Total trainers
    cursor.execute("SELECT COUNT(*) AS total FROM trainers")
    total_trainers = cursor.fetchone()["total"]

    # Available seats
    cursor.execute("SELECT COALESCE(SUM(seats), 0) AS total FROM courses")
    total_seats = cursor.fetchone()["total"]

    # Total sales
    cursor.execute("SELECT COALESCE(SUM(amount), 0) AS total FROM sales")
    total_sales = cursor.fetchone()["total"]

    # Courses
    cursor.execute("""SELECT courses.id, courses.name, courses.price,courses.seats, courses.trainer
        FROM courses ORDER BY courses.id DESC LIMIT 5""")
    courses = cursor.fetchall()

    # Low stock courses
    cursor.execute("""SELECT id, name, seats FROM courses WHERE seats <= 5 ORDER BY seats ASC""")
    low_stock_courses = cursor.fetchall()
    cursor.close()
    connection.close()

    return render_template(
        "dashboard.html",
        username=session["username"],
        total_courses=total_courses,
        total_trainers=total_trainers,
        total_seats=total_seats,
        total_sales=total_sales,
        courses=courses,
        low_stock_courses=low_stock_courses
    )

import os
import barcode
from barcode.writer import ImageWriter

@app.route("/generate-barcode/<int:course_id>")
def generate_barcode(course_id):

    if "username" not in session:
        return redirect(url_for("login"))

    barcode_folder = os.path.join(app.static_folder, "barcodes")
    os.makedirs(barcode_folder, exist_ok=True)
    filename = os.path.join(barcode_folder, f"course_{course_id}")
    try:
        scan_url = url_for("scan_course", course_id=course_id, _external=True)
        print("Barcode URL:", scan_url)
        code128 = barcode.get("code128", str(course_id), writer=ImageWriter())
        code128.save(filename, options={"module_width": 0.5, "module_height": 30, "quiet_zone": 10,
                "font_size": 14, "text_distance": 8, "write_text": True, "dpi": 300}
        )
        flash("Barcode generated successfully.", "success")

    except Exception as e:
        print("Barcode Error:", e)
        flash("Unable to generate barcode.", "error")

    return redirect(url_for("courses"))

@app.route("/scan-course/<int:course_id>")
def scan_course(course_id):
    connection = get_db_connection()

    if connection is None:
        return "Database connection failed.", 500

    cursor = connection.cursor(dictionary=True)
    cursor.execute("""SELECT id, name, category, trainer, price, seats FROM courses WHERE id = %s""", (course_id,))
    course = cursor.fetchone()
    cursor.close()
    connection.close()

    if not course:
        return """<h2 style="text-align:center; margin-top:50px;">Course not found</h2>""", 404
    return render_template("scan_course.html", course=course)

@app.route("/scan-barcode")
def scan_barcode():

    if "admin_id" not in session:
        return redirect(url_for("login"))
    return render_template("scan_barcode.html")

# COURSES
@app.route("/courses")
def courses():

    if "username" not in session:
        return redirect(url_for("login"))

    connection = get_db_connection()
    cursor = connection.cursor(dictionary=True)
    cursor.execute("""SELECT courses.id, courses.name, courses.category, courses.price, courses.seats, courses.trainer FROM courses ORDER BY courses.id DESC""")
    courses = cursor.fetchall()
    cursor.close()
    connection.close()
    return render_template("courses.html",username=session["username"],courses=courses)

# ADD COURSE
@app.route('/add-course', methods=['GET', 'POST'])
def add_course():

    if request.method == 'POST':
        course_name = request.form['course_name']
        category = request.form['category']
        trainer = request.form['trainer']
        price = request.form['price']
        seats = request.form['seats']

        return redirect(url_for('courses'))
    return render_template('add_course.html')

@app.route('/edit-course/<int:course_id>', methods=['GET', 'POST'])
def edit_course(course_id):

    if "admin_id" not in session:
        return redirect(url_for("login"))

    connection = get_db_connection()

    if connection is None:
        flash("Database connection failed.", "error")
        return redirect(url_for("courses"))

    cursor = connection.cursor(dictionary=True)

    if request.method == 'POST':

        course_name = request.form.get('course_name')
        category = request.form.get('category')
        trainer_id = request.form.get('trainer')
        price = request.form.get('price')
        seats = request.form.get('seats')

        # Get trainer NAME from trainer ID
        cursor.execute("SELECT name FROM trainers WHERE id = %s", (trainer_id,))
        trainer_data = cursor.fetchone()
        trainer_name = trainer_data["name"] if trainer_data else None

        # Update course
        query = """UPDATE courses SET name = %s, category = %s, trainer = %s, price = %s, seats = %s WHERE id = %s"""
        cursor.execute(query,(course_name, category, trainer_name, price, seats, course_id))
        connection.commit()
        cursor.close()
        connection.close()

        flash("Course updated successfully.", "success")
        return redirect(url_for("courses"))

    # Get existing course
    cursor.execute("""SELECT id, name, category, trainer, price, seats FROM courses WHERE id = %s""", (course_id,))
    course = cursor.fetchone()

    # Get all trainers
    cursor.execute("""SELECT id, name FROM trainers ORDER BY name""")
    trainers = cursor.fetchall()
    cursor.close()
    connection.close()

    if course is None:
        return redirect(url_for("courses"))

    return render_template("edit_course.html", course=course, trainers=trainers, username=session["username"])

# DELETE COURSE
@app.route("/delete-course/<int:course_id>")
def delete_course(course_id):

    if "admin_id" not in session:
        return redirect(url_for("login"))

    connection = get_db_connection()
    cursor = connection.cursor()

    try:
        cursor.execute("DELETE FROM courses WHERE id = %s",(course_id,))
        connection.commit()

    finally:
        cursor.close()
        connection.close()

    return redirect(url_for("courses"))

# TRAINERS
@app.route("/trainers")
def trainers():
    if "admin_id" not in session:
        return redirect(url_for("login"))

    connection = get_db_connection()
    cursor = connection.cursor(dictionary=True)
    cursor.execute("""SELECT id, name, email, phone, specialization FROM trainers ORDER BY id DESC""")
    trainers = cursor.fetchall()
    cursor.close()
    connection.close()
    return render_template("trainers.html", username=session["username"], trainers=trainers)

# ADD TRAINER
@app.route('/add-trainer', methods=['GET', 'POST'])
def add_trainer():

    if "admin_id" not in session:
        return redirect(url_for("login"))

    if request.method == 'POST':
        trainer_name = request.form.get('name')
        specialization = request.form.get('specialization')
        email = request.form.get('email')
        phone = request.form.get('phone')
        connection = get_db_connection()

        if connection is None:
            flash("Database connection failed.", "error")
            return redirect(url_for("add_trainer"))

        cursor = connection.cursor()
        query = """INSERT INTO trainers (name, email, phone, specialization) VALUES (%s, %s, %s, %s)"""
        cursor.execute(query, (trainer_name, email, phone, specialization))
        connection.commit()
        cursor.close()
        connection.close()
        return redirect(url_for('trainers'))
    return render_template('add_trainer.html')

# DELETE TRAINER
@app.route('/delete-trainer/<int:trainer_id>', methods=['POST'])
def delete_trainer(trainer_id):

    if "admin_id" not in session:
        return redirect(url_for("login"))

    connection = get_db_connection()

    if connection is None:
        flash("Database connection failed.", "error")
        return redirect(url_for("trainers"))

    cursor = connection.cursor()

    try:
        cursor.execute("SELECT name FROM trainers WHERE id = %s", (trainer_id,))
        trainer_data = cursor.fetchone()

        if trainer_data:
            trainer_name = trainer_data[0]
            cursor.execute("UPDATE courses SET trainer = NULL WHERE trainer = %s", (trainer_name,))

        cursor.execute("DELETE FROM trainers WHERE id = %s", (trainer_id,))
        connection.commit()
        flash("Trainer deleted successfully.", "success")

    except Exception as e:
        connection.rollback()
        flash("Unable to delete trainer.", "error")

    finally:
        cursor.close()
        connection.close()

    return redirect(url_for("trainers"))

# SALES
@app.route("/sales")
def sales():
    if "admin_id" not in session:
        return redirect(url_for("login"))

    connection = get_db_connection()
    cursor = connection.cursor(dictionary=True)
    cursor.execute("""SELECT s.id, s.student_name, c.name AS course, s.amount, s.sale_date
        FROM sales s JOIN courses c ON s.course_id = c.id ORDER BY s.id DESC""")
    sales_data = cursor.fetchall()
    cursor.close()
    connection.close()
    return render_template("sales.html",username=session["username"],sales=sales_data)

@app.route("/add-sale", methods=["GET", "POST"])
def add_sale():
    if "admin_id" not in session:
        return redirect(url_for("login"))

    connection = get_db_connection()
    cursor = connection.cursor(dictionary=True)

    if request.method == "POST":
        student_name = request.form.get("student_name")
        course_id = request.form.get("course_id")
        amount = request.form.get("amount")
        sale_date = request.form.get("sale_date")

        cursor.execute("""INSERT INTO sales (student_name, course_id, amount, sale_date) VALUES (%s, %s, %s, %s)""", (student_name, course_id, amount, sale_date))
        connection.commit()
        cursor.close()
        connection.close()
        return redirect(url_for("sales"))

    cursor.execute("SELECT id, name, price FROM courses ORDER BY name")
    courses = cursor.fetchall()
    cursor.close()
    connection.close()
    return render_template("add_sale.html",username=session["username"], courses=courses)

@app.route("/sales-report")
def sales_report():

    if "admin_id" not in session:
        return redirect(url_for("login"))

    connection = get_db_connection()
    cursor = connection.cursor(dictionary=True)

    # Total sales, total enrollments and average sale
    cursor.execute("""SELECT COALESCE(SUM(amount), 0) AS total_sales, COUNT(*) AS total_enrollments, COALESCE(AVG(amount), 0) AS average_sale FROM sales""")
    summary = cursor.fetchone()

    # Sales summary by course
    cursor.execute("""SELECT c.name AS course, COUNT(s.id) AS enrollments, COALESCE(SUM(s.amount), 0) AS total_revenue 
        FROM sales s JOIN courses c ON s.course_id = c.id GROUP BY c.id, c.name ORDER BY total_revenue DESC""")

    sales_summary = cursor.fetchall()
    cursor.close()
    connection.close()

    return render_template("sales_report.html",username=session["username"],total_sales=summary["total_sales"], 
            total_enrollments=summary["total_enrollments"], average_sale=summary["average_sale"], sales_summary=sales_summary
    )

@app.route("/download-sales-excel")
def download_sales_excel():
    if "admin_id" not in session:
        return redirect(url_for("login"))

    connection = get_db_connection()
    cursor = connection.cursor(dictionary=True)
    cursor.execute("""SELECT id, student_name, course_id, amount, sale_date FROM sales ORDER BY sale_date DESC""")
    sales_data = cursor.fetchall()
    cursor.close()
    connection.close()

    import openpyxl
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    from io import BytesIO
    from flask import send_file

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Sales Report"
    headers = ["Sale ID", "Student Name", "Course", "Amount", "Sale Date"]

    for col, header in enumerate(headers, 1):
        cell = worksheet.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for row_num, sale in enumerate(sales_data, 2):
        worksheet.cell(row=row_num, column=1, value=sale["id"])
        worksheet.cell(row=row_num, column=2, value=sale["student_name"])
        worksheet.cell(row=row_num, column=3, value=sale["course_id"])
        worksheet.cell(row=row_num, column=4, value=float(sale["amount"]))
        worksheet.cell(row=row_num, column=5, value=str(sale["sale_date"]))

    worksheet.column_dimensions["A"].width = 12
    worksheet.column_dimensions["B"].width = 25
    worksheet.column_dimensions["C"].width = 25
    worksheet.column_dimensions["D"].width = 15
    worksheet.column_dimensions["E"].width = 20

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name="NetTech_Sales_Report.xlsx", 
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

# LOGOUT
@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))

# RUN
if __name__ == "__main__":
    app.run(debug=True)